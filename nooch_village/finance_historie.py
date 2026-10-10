"""Historie per maand — Finance PR 2a (10 oktober 2026).

Besluit Stefan: voor maanden zonder kostenregister (2025, begin 2026) vult hij de bedragen per
maand per regel in, met de cijfers van Venema of eigen cijfers, met de hand of via een eenmalige
import. Het P&L combineert die historie met het register zonder dubbel te tellen.

DRIE REGELS DIE JE NIET MAG OMDRAAIEN:
1. ALLEEN TOEVOEGEN. `data/finance/history.jsonl` groeit. Een wijziging is een nieuwe regel met
   `supersedes` = de vorige waarde van dezelfde cel; de actuele waarde van een cel is de regel die
   door niets is vervangen. Niets wordt overschreven of verwijderd, en per cel is de geschiedenis
   terug te lezen. Een import staat in `history_imports.jsonl` (ook alleen toevoegen: een nieuwe
   status is een nieuwe regel), met het bestand zelf onder `data/finance/imports/` en zijn sha256.
2. CENTEN, EXACT. Euro's met hoogstens twee decimalen gaan via `Decimal` naar centen; meer
   decimalen is een fout, geen afronding (behalve een XLSX-getal, dat Excel als float bewaart).
3. ÉÉN BRON PER MAAND. `history_until` (YYYY-MM, leeg = geen historie): maanden tot en met die
   maand komen uit de historie, daarna uit het register. Registerregels vóór de cutover blijven
   staan maar tellen niet (`kies_bronnen`); het scherm meldt hoeveel het er zijn. De instelling
   staat in `history_settings.jsonl`: elke wijziging een regel (wie, wanneer, oud, nieuw), de
   laatste is de waarde.

Een lege cel is GEEN DATA, geen nul. Een cel leegmaken kan (nog) niet: dat zou een regel zonder
bedrag vragen, en dat is een andere datastructuur dan besloten — het scherm zegt dat erbij.

`net_sales_pairs` bestaat alleen in het raster en in het bestand: de paren zijn een veld van de
net_sales-regel van die maand (`pairs`), zoals in het opbrengstregister.

FINANCE-DATA BLIJFT HIER: zelfde guard als het kosten- en opbrengstregister.
"""
from __future__ import annotations

import csv
import io
import os
import re
import time
import uuid
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from nooch_village.finance_register import _PERIODE, MAP, _Register, _tekst, sha256

BRON = "history"
LABEL = "history (manual or imported), not backed by documents"
NET_SALES = "net_sales"
PAREN = "net_sales_pairs"                       # alleen raster/bestand; opgeslagen als `pairs`
IMPORT_EXT = {"csv": "text/csv", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
MAX_IMPORT = 10 * 1024 * 1024


class HistorieRegister(_Register):
    """`data/finance/history.jsonl`. Alleen toevoegen — zie `_Register`."""

    BESTAND = "history.jsonl"
    BEDRAGEN = ()

    def actueel(self) -> dict[tuple[str, str], dict]:
        """{(period, line_key): regel} — per cel de regel die door niets is vervangen (bij twee
        onvervangen regels voor één cel, wat alleen bij een gelijktijdige opslag kan: de laatste)."""
        regels = self.alle()
        vervangen = {r.get("supersedes") for r in regels if r.get("supersedes")}
        uit: dict[tuple[str, str], dict] = {}
        for r in sorted(regels, key=lambda x: x.get("created_at", 0)):
            if r["id"] not in vervangen:
                uit[(r["period"], r["line_key"])] = r
        return uit

    def geschiedenis(self, period: str, line_key: str) -> list[dict]:
        """Alle waarden die deze cel ooit had, nieuwste eerst."""
        return sorted((r for r in self.alle() if r.get("period") == period and r.get("line_key") == line_key),
                      key=lambda x: x.get("created_at", 0), reverse=True)

    def schrijf(self, wijzigingen: list[dict], *, door: str, bron: str = "manual", note: str = "") -> list[dict]:
        """Eén append voor alle gewijzigde cellen van één opslag of import."""
        regels = [self._stempel({"period": w["period"], "line_key": w["line_key"],
                                 "amount_cents": w["amount_cents"], "pairs": w.get("pairs"),
                                 "note": _tekst(note, 500), "source": bron,
                                 "supersedes": w.get("supersedes") or ""}, door) for w in wijzigingen]
        if regels:
            self._append(regels)
        return regels


class ImportLog(_Register):
    """`data/finance/history_imports.jsonl` plus de bestanden in `data/finance/imports/`."""

    BESTAND = "history_imports.jsonl"
    BEDRAGEN = ()

    def bewaar(self, bestandsnaam: str, blob: bytes, *, door: str) -> tuple[dict | None, str]:
        """Bewaar het geuploade bestand (nooit overschrijven) en log de import als `previewed`."""
        ext = os.path.splitext(bestandsnaam or "")[1].lower().lstrip(".")
        if ext not in IMPORT_EXT:
            return None, "only csv or xlsx"
        if not blob:
            return None, "the file is empty"
        if len(blob) > MAX_IMPORT:
            return None, "the file is larger than 10 MB"
        if ext == "xlsx" and not blob.startswith(b"PK\x03\x04"):
            return None, "this file is not a real .xlsx"
        rel = os.path.join(MAP, "imports", f"{uuid.uuid4().hex}.{ext}")
        vol = os.path.join(os.path.dirname(self.dir), rel)
        os.makedirs(os.path.dirname(vol), exist_ok=True)
        with open(vol, "xb") as f:
            f.write(blob)
        r = self._stempel({"import_id": uuid.uuid4().hex[:12], "status": "previewed",
                           "attachment": {"path": rel, "sha256": sha256(blob), "size": len(blob),
                                          "filename": os.path.basename(bestandsnaam)[:120],
                                          "mime": IMPORT_EXT[ext]}}, door)
        self._append([r])
        return r, ""

    def status(self, import_id: str) -> dict | None:
        """De laatste regel van deze import (de actuele status)."""
        eigen = [r for r in self.alle() if r.get("import_id") == import_id]
        return max(eigen, key=lambda x: x.get("created_at", 0)) if eigen else None

    def meld(self, basis: dict, *, door: str, **velden) -> dict:
        r = self._stempel({k: v for k, v in basis.items() if k not in ("id", "created_at", "created_by")}
                          | velden, door)
        self._append([r])
        return r


class Instellingen(_Register):
    """`data/finance/history_settings.jsonl`: de cutover. Elke wijziging een regel; de laatste telt."""

    BESTAND = "history_settings.jsonl"
    BEDRAGEN = ()

    def history_until(self) -> str:
        eigen = [r for r in self.alle() if r.get("key") == "history_until"]
        return max(eigen, key=lambda x: x.get("created_at", 0)).get("new", "") if eigen else ""

    def zet_history_until(self, waarde: str, *, door: str) -> tuple[bool, str]:
        waarde = (waarde or "").strip()
        if waarde and not _PERIODE.match(waarde):
            return False, "the cutover must be YYYY-MM, or empty for no history"
        oud = self.history_until()
        if oud == waarde:
            return True, "unchanged"
        self.voeg_toe({"key": "history_until", "old": oud, "new": waarde}, door=door)
        return True, ""


# ── regels in P&L-volgorde ──────────────────────────────────────────────────

def lijnen(cfg: dict) -> list[dict]:
    """De rijen van het raster en het bestand, in P&L-volgorde: {key, label, soort} met soort
    `bedrag`, `paren` of `kop` (een groepskop, geen invoer)."""
    rev = list(cfg.get("revenue_lines") or [])
    uit = []
    for r in rev:
        uit.append({"key": r["key"], "label": r["label"], "soort": "bedrag"})
        if r["key"] == NET_SALES:
            uit.append({"key": PAREN, "label": "Net sales pairs", "soort": "paren"})
    for g in cfg["management_lines"]:
        if g.get("label"):
            uit.append({"key": g["groep"], "label": g["label"], "soort": "kop"})
        for r in g["regels"]:
            uit.append({"key": r["key"], "label": r["label"], "soort": "bedrag",
                        "negatief": bool(r.get("negatief"))})
    return uit


def _sleutels(cfg: dict) -> dict[str, dict]:
    return {x["key"]: x for x in lijnen(cfg) if x["soort"] != "kop"}


# ── bedragen ────────────────────────────────────────────────────────────────

_DUIZEND_PUNT = re.compile(r"^-?\d{1,3}(\.\d{3})+$")
_DUIZEND_KOMMA = re.compile(r"^-?\d{1,3}(,\d{3})+$")


def bedrag(x) -> tuple[int | None, str]:
    """Euro's → (centen, "") of (None, reden). Herkent decimale komma én punt, duizendtallen
    (1.234,56 en 1,234.56) en een €-teken. Meer dan twee decimalen is een fout: geen stille afronding.
    Een Decimal of float (uit XLSX) wordt op centen afgerond — Excel bewaart getallen als float."""
    if isinstance(x, (int, float, Decimal)) and not isinstance(x, bool):
        d = Decimal(str(x))
        return int((d * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP)), ""
    s = str(x if x is not None else "").strip().replace("€", "").replace(" ", "").replace(" ", "")
    if not s:
        return None, ""
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(",", ".")
    elif _DUIZEND_PUNT.match(s):
        s = s.replace(".", "")
    elif _DUIZEND_KOMMA.match(s) and len(s.split(",")) > 2:
        s = s.replace(",", "")                     # 1,234,567 — meer dan één groep: duizendtallen
    else:
        s = s.replace(",", ".")
    try:
        d = Decimal(s)
    except InvalidOperation:
        return None, f"not an amount: {x!r}"
    if not d.is_finite():
        return None, f"not an amount: {x!r}"
    if d != d.quantize(Decimal("0.01")):
        return None, f"more than two decimals: {x!r}"
    return int(d * 100), ""


def paren(x) -> tuple[int | None, str]:
    if isinstance(x, (int, float, Decimal)) and not isinstance(x, bool):
        if Decimal(str(x)) != Decimal(str(x)).to_integral_value():
            return None, f"pairs must be a whole number: {x!r}"
        n = int(Decimal(str(x)))
    else:
        s = str(x if x is not None else "").strip().replace(".", "").replace(",", "")
        if not s:
            return None, ""
        try:
            n = int(s)
        except ValueError:
            return None, f"pairs must be a whole number: {x!r}"
    if n < 0:
        return None, "pairs cannot be negative"
    return n, ""


def euro_nl(c: int) -> str:
    teken = "-" if c < 0 else ""
    c = abs(int(c))
    return f"{teken}{c // 100},{c % 100:02d}"


# ── vergelijken: wat verandert er ───────────────────────────────────────────

def vergelijk(bedragen: dict[tuple[str, str], int], paren_: dict[str, int],
              actueel: dict[tuple[str, str], dict], cfg: dict) -> dict:
    """Gewenste celwaarden tegen de actuele → {nieuw, gewijzigd, ongewijzigd, fouten, schrijf}.
    `schrijf` zijn de regels die een opslag of import zou toevoegen (met `supersedes`)."""
    sleutels = _sleutels(cfg)
    uit = {"nieuw": 0, "gewijzigd": 0, "ongewijzigd": 0, "fouten": [], "schrijf": []}
    cellen = set(bedragen) | {(p, NET_SALES) for p in paren_}
    for (p, k) in sorted(cellen):
        lijn = sleutels.get(k)
        if lijn is None or not _PERIODE.match(p):
            uit["fouten"].append(f"{p} {k}: unknown line or month")
            continue
        oud = actueel.get((p, k))
        cents = bedragen.get((p, k), oud["amount_cents"] if oud else None)
        if cents is None:
            uit["fouten"].append(f"{p}: pairs without net sales — enter the net sales of that month too")
            continue
        if cents < 0 and not lijn.get("negatief"):
            uit["fouten"].append(f"{p} {lijn['label']}: negative amounts only on lines that allow a credit")
            continue
        nieuw_paren = paren_.get(p, oud.get("pairs") if oud else None) if k == NET_SALES else None
        if oud is None:
            uit["nieuw"] += 1
        elif oud["amount_cents"] == cents and oud.get("pairs") == nieuw_paren:
            uit["ongewijzigd"] += 1
            continue
        else:
            uit["gewijzigd"] += 1
        uit["schrijf"].append({"period": p, "line_key": k, "amount_cents": cents, "pairs": nieuw_paren,
                               "supersedes": oud["id"] if oud else "", "oud": oud})
    return uit


# ── bestanden lezen ─────────────────────────────────────────────────────────

def _rijen_csv(blob: bytes) -> list[list]:
    try:
        tekst = blob.decode("utf-8-sig")
    except UnicodeDecodeError:
        tekst = blob.decode("cp1252")
    kop = tekst.splitlines()[0] if tekst.strip() else ""
    scheiding = max((";", ",", "\t"), key=kop.count)
    return [r for r in csv.reader(io.StringIO(tekst), delimiter=scheiding)]


def _rijen_xlsx(blob: bytes) -> list[list]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        return [["" if v is None else v for v in rij] for rij in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def _maand(x) -> str:
    if hasattr(x, "strftime"):                       # Excel maakt van "2025-01" soms een datum
        return x.strftime("%Y-%m")
    s = str(x or "").strip()
    return s if _PERIODE.match(s) else ""


def lees_bestand(bestandsnaam: str, blob: bytes, cfg: dict) -> dict:
    """Een export-formaat bestand → {bedragen, paren, fouten, rijen}. Regels op line_key, anders op
    het Engelse label; een onbekende regel of maand is een fout (met rijnummer), nooit stil weg."""
    ext = os.path.splitext(bestandsnaam or "")[1].lower().lstrip(".")
    try:
        rijen = _rijen_xlsx(blob) if ext == "xlsx" else _rijen_csv(blob)
    except Exception as e:                          # een kapot bestand is een fout, geen crash
        return {"bedragen": {}, "paren": {}, "fouten": [f"the file cannot be read: {type(e).__name__}"],
                "rijen": 0}
    uit = {"bedragen": {}, "paren": {}, "fouten": [], "rijen": 0}
    if not rijen:
        uit["fouten"].append("the file is empty")
        return uit
    kop = rijen[0]
    maanden = {}
    for i, h in enumerate(kop[2:], start=2):
        m = _maand(h)
        if m:
            maanden[i] = m
        elif str(h or "").strip():
            uit["fouten"].append(f"column {i + 1}: '{h}' is not a month (YYYY-MM)")
    sleutels = _sleutels(cfg)
    op_label = {x["label"].strip().lower(): x["key"] for x in sleutels.values()}
    for nr, rij in enumerate(rijen[1:], start=2):
        if not any(str(v).strip() for v in rij):
            continue
        uit["rijen"] += 1
        key = str(rij[0] if rij else "").strip()
        label = str(rij[1] if len(rij) > 1 else "").strip()
        if key not in sleutels:
            key = op_label.get(label.lower()) or op_label.get(key.lower()) or ""
        if not key:
            uit["fouten"].append(f"row {nr}: unknown line '{rij[0] if rij else ''}' / '{label}'")
            continue
        for i, m in maanden.items():
            v = rij[i] if i < len(rij) else ""
            if key == PAREN:
                n, fout = paren(v)
                if fout:
                    uit["fouten"].append(f"row {nr}, {m}: {fout}")
                elif n is not None:
                    uit["paren"][m] = n
            else:
                c, fout = bedrag(v)
                if fout:
                    uit["fouten"].append(f"row {nr}, {m}: {fout}")
                elif c is not None:
                    uit["bedragen"][(m, key)] = c
    return uit


# ── exporteren ──────────────────────────────────────────────────────────────

def maanden_van(jaar: str) -> list[str]:
    return [f"{jaar}-{m:02d}" for m in range(1, 13)]


def export_tabel(actueel: dict[tuple[str, str], dict], cfg: dict, maanden: list[str]) -> list[list]:
    """[kop, rijen...] met centen (int) of None per maand — het sjabloon én de export."""
    tabel = [["line_key", "label", *maanden]]
    for x in lijnen(cfg):
        if x["soort"] == "kop":
            continue
        rij = [x["key"], x["label"]]
        for m in maanden:
            if x["soort"] == "paren":
                r = actueel.get((m, NET_SALES))
                rij.append(r.get("pairs") if r else None)
            else:
                r = actueel.get((m, x["key"]))
                rij.append(r["amount_cents"] if r else None)
        tabel.append(rij)
    return tabel


def export_csv(tabel: list[list]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(tabel[0])
    for rij in tabel[1:]:
        paar = rij[0] == PAREN
        w.writerow(rij[:2] + ["" if v is None else (str(v) if paar else euro_nl(v)) for v in rij[2:]])
    return buf.getvalue()


def export_xlsx(tabel: list[list]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "history"
    ws.append([str(v) for v in tabel[0]])           # maanden als TEKST, anders maakt Excel er data van
    for rij in tabel[1:]:
        paar = rij[0] == PAREN
        ws.append(rij[:2] + [None if v is None else (v if paar else Decimal(v) / 100) for v in rij[2:]])
        if not paar:
            for cel in ws[ws.max_row][2:]:
                cel.number_format = "0.00"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── het P&L: welke bron per maand ───────────────────────────────────────────

def posten(actueel: dict[tuple[str, str], dict], cfg: dict) -> tuple[list[tuple], list[tuple], dict[str, int]]:
    """(kostenposten, opbrengstposten, paren) uit de historie — zelfde vorm als de registers."""
    rev = {r["key"] for r in cfg.get("revenue_lines") or []}
    kosten, opbr, par = [], [], {}
    for (p, k), r in actueel.items():
        (opbr if k in rev else kosten).append((p, k, int(r["amount_cents"]), BRON))
        if k == NET_SALES and isinstance(r.get("pairs"), int):
            par[p] = r["pairs"]
    return kosten, opbr, par


def kies_bronnen(register_kosten: list[tuple], register_opbr: list[tuple], hist_kosten: list[tuple],
                 hist_opbr: list[tuple], tot: str) -> tuple[list[tuple], list[tuple], int]:
    """Eén bron per maand: tot en met `tot` de historie, daarna het register. → (kosten,
    opbrengsten, aantal registerposten vóór de cutover dat NIET meetelt). Leeg `tot` = geen historie."""
    if not tot:
        return list(register_kosten), list(register_opbr), 0
    genegeerd = sum(1 for p in list(register_kosten) + list(register_opbr) if p[0] <= tot)
    return ([p for p in register_kosten if p[0] > tot] + [p for p in hist_kosten if p[0] <= tot],
            [p for p in register_opbr if p[0] > tot] + [p for p in hist_opbr if p[0] <= tot],
            genegeerd)


def nu_maand() -> str:
    return time.strftime("%Y-%m")
